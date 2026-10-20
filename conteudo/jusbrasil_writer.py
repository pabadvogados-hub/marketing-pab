"""
JUSBRASIL Writer - orquestrador.
Busca noticias, gera artigo SEO e salva DOCX no Drive para revisao humana.

Uso:
    py CONTEUDO/jusbrasil_writer.py                  # fluxo completo
    py CONTEUDO/jusbrasil_writer.py --sem-notificar  # pula WhatsApp
    py CONTEUDO/jusbrasil_writer.py --teste          # nao salva no Drive
    py CONTEUDO/jusbrasil_writer.py --dias 14        # janela de busca
"""
import argparse
import io
import json
import os
import re
import sys
from datetime import datetime

PROJECT_ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, PROJECT_ROOT)
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

from dotenv import load_dotenv
load_dotenv(os.path.join(PROJECT_ROOT, "config", ".env"))

from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from googleapiclient.http import MediaIoBaseUpload

from INTEGRACOES.google_integration import autenticar_google
from INTEGRACOES.atendedireito_integration import (
    buscar_contato_por_telefone,
    enviar_mensagem_texto,
)

from buscar_noticias import buscar_noticias_caminhoneiro
from gerar_artigo import gerar_artigo, escolher_melhor_noticia


HISTORICO_PATH = os.path.join(os.path.dirname(os.path.abspath(__file__)), "historico.json")


def carregar_historico() -> list:
    if not os.path.exists(HISTORICO_PATH):
        return []
    try:
        with open(HISTORICO_PATH, "r", encoding="utf-8") as f:
            return json.load(f)
    except Exception:
        return []


def salvar_historico(historico: list):
    with open(HISTORICO_PATH, "w", encoding="utf-8") as f:
        json.dump(historico, f, ensure_ascii=False, indent=2)


def _render_inline(p, linha: str):
    """Processa **bold** e *italic* inline dentro de um paragrafo."""
    partes = re.split(r"(\*\*[^*]+\*\*|\*[^*\s][^*]*[^*\s]\*)", linha)
    for parte in partes:
        if not parte:
            continue
        if parte.startswith("**") and parte.endswith("**"):
            r = p.add_run(parte[2:-2])
            r.bold = True
        elif parte.startswith("*") and parte.endswith("*") and len(parte) > 2:
            r = p.add_run(parte[1:-1])
            r.italic = True
        else:
            p.add_run(parte)


def markdown_para_docx(artigo: dict) -> bytes:
    """Converte o artigo em DOCX formatado com Montserrat 11pt."""
    doc = Document()

    # Estilo base
    estilo = doc.styles["Normal"]
    estilo.font.name = "Montserrat"
    estilo.font.size = Pt(11)

    # Bloco de metadados para revisao
    meta = doc.add_paragraph()
    meta_run = meta.add_run(
        "[METADADOS PARA REVISAO - APAGAR ANTES DE PUBLICAR NO JUSBRASIL]\n"
        f"Titulo SEO: {artigo.get('titulo', '')}\n"
        f"Meta description: {artigo.get('meta_description', '')}\n"
        f"Palavra-chave principal: {artigo.get('palavra_chave_principal', '')}\n"
        f"Tags: {', '.join(artigo.get('palavras_chave_secundarias', []))}\n\n"
        f"Observacoes: {artigo.get('observacoes_revisao', '')}\n"
        f"Gancho: {artigo.get('_noticia_origem', {}).get('titulo', '')}\n"
        f"Fonte: {artigo.get('_noticia_origem', {}).get('link', '')}\n"
        f"{'-' * 70}"
    )
    meta_run.font.size = Pt(9)
    meta_run.font.color.rgb = RGBColor(0x80, 0x80, 0x80)
    doc.add_paragraph()

    # Corpo
    corpo = artigo.get("corpo_markdown", "").strip()
    for linha in corpo.split("\n"):
        linha = linha.rstrip()
        if not linha:
            doc.add_paragraph()
            continue

        if linha.startswith("# "):
            p = doc.add_paragraph()
            r = p.add_run(linha[2:].strip())
            r.bold = True
            r.font.size = Pt(16)
        elif linha.startswith("## "):
            p = doc.add_paragraph()
            r = p.add_run(linha[3:].strip())
            r.bold = True
            r.font.size = Pt(13)
        elif linha.startswith("### "):
            p = doc.add_paragraph()
            r = p.add_run(linha[4:].strip())
            r.bold = True
            r.font.size = Pt(11)
        elif linha.startswith("- ") or linha.startswith("* "):
            p = doc.add_paragraph(style="List Bullet")
            _render_inline(p, linha[2:].strip())
        elif re.match(r"^\d+\.\s", linha):
            p = doc.add_paragraph(style="List Number")
            _render_inline(p, re.sub(r"^\d+\.\s", "", linha))
        elif linha.startswith("> "):
            p = doc.add_paragraph()
            r = p.add_run(linha[2:].strip())
            r.italic = True
            p.paragraph_format.left_indent = Cm(2)
        else:
            p = doc.add_paragraph()
            _render_inline(p, linha)
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    buffer = io.BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer.read()


def salvar_no_drive(drive_service, docx_bytes: bytes, nome_arquivo: str, pasta_id: str):
    media = MediaIoBaseUpload(
        io.BytesIO(docx_bytes),
        mimetype="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        resumable=False,
    )
    metadata = {"name": nome_arquivo}
    if pasta_id:
        metadata["parents"] = [pasta_id]
    arquivo = drive_service.files().create(
        body=metadata, media_body=media, fields="id, webViewLink",
    ).execute()
    return arquivo["id"], arquivo.get("webViewLink", "")


def notificar_whatsapp(artigo: dict, link_drive: str) -> bool:
    telefone = os.getenv("JUSBRASIL_NOTIFICAR_TELEFONE", "").strip()
    if not telefone:
        print("   JUSBRASIL_NOTIFICAR_TELEFONE nao configurado no .env - pulando")
        return False

    contato = buscar_contato_por_telefone(telefone)
    if not contato:
        print(f"   Contato {telefone} nao encontrado no Atende Direito")
        return False

    mensagem = (
        "*Artigo JUSBRASIL pronto para revisao*\n\n"
        f"*Titulo:* {artigo.get('titulo', '')}\n\n"
        f"*Palavra-chave:* {artigo.get('palavra_chave_principal', '')}\n\n"
        f"*Resumo:* {artigo.get('resumo_executivo', '')}\n\n"
        f"*A revisar:* {artigo.get('observacoes_revisao', '')}\n\n"
        f"*Draft no Drive:* {link_drive}\n\n"
        "Fluxo:\n"
        "1) Abra o DOCX e revise\n"
        "2) Apague o bloco de metadados no topo\n"
        "3) Cole o corpo no editor do JUSBRASIL\n"
        "4) Use titulo e meta description ao publicar\n\n"
        "PAB Automacoes"
    )
    user_ns = contato.get("user_ns", "")
    return enviar_mensagem_texto(user_ns, mensagem)


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--sem-notificar", action="store_true", help="nao envia WhatsApp")
    parser.add_argument("--teste", action="store_true", help="gera mas nao salva no Drive")
    parser.add_argument("--dias", type=int, default=7, help="janela de busca (default 7)")
    args = parser.parse_args()

    print("=" * 64)
    print(f"JUSBRASIL WRITER - {datetime.now().strftime('%d/%m/%Y %H:%M')}")
    print("=" * 64)

    # 1. Buscar noticias
    print("\n[1/5] Buscando noticias recentes...")
    noticias = buscar_noticias_caminhoneiro(dias=args.dias)
    if not noticias:
        print("Nenhuma noticia encontrada. Abortando.")
        return 1

    # 2. Filtrar historico
    print("\n[2/5] Filtrando contra historico...")
    historico = carregar_historico()
    links_usados = {h.get("link_origem") for h in historico}
    titulos_usados = {(h.get("titulo_origem") or "").lower()[:60] for h in historico}
    noticias_inedit = [
        n for n in noticias
        if n["link"] not in links_usados
        and n["titulo"].lower()[:60] not in titulos_usados
    ]
    print(f"   {len(noticias_inedit)} ineditas (de {len(noticias)} totais)")
    if not noticias_inedit:
        print("Todas as noticias ja foram usadas. Tente --dias maior.")
        return 1

    # 3. Escolher melhor
    print("\n[3/5] Escolhendo melhor tema (Haiku)...")
    noticia = escolher_melhor_noticia(noticias_inedit, historico)
    print(f"   Escolhida: [{noticia['fonte']}] {noticia['titulo']}")

    # 4. Gerar artigo
    print("\n[4/5] Gerando artigo com Claude Opus 4.6...")
    artigo = gerar_artigo(noticia)
    n_palavras = len(artigo.get("corpo_markdown", "").split())
    print(f"   Titulo: {artigo.get('titulo', '')}")
    print(f"   Palavras: ~{n_palavras}")
    tk = artigo.get("_tokens", {})
    print(f"   Tokens: in={tk.get('input')} out={tk.get('output')} cache_r={tk.get('cache_read')}")

    if args.teste:
        out = os.path.join(os.path.dirname(os.path.abspath(__file__)), "_teste_artigo.json")
        with open(out, "w", encoding="utf-8") as f:
            json.dump(artigo, f, ensure_ascii=False, indent=2)
        print(f"\n[TESTE] Artigo salvo em: {out}")
        return 0

    # 5. Salvar no Drive e notificar
    print("\n[5/5] Salvando DOCX no Drive...")
    drive_service, _ = autenticar_google()
    docx_bytes = markdown_para_docx(artigo)
    data_str = datetime.now().strftime("%Y-%m-%d")
    titulo_limpo = re.sub(r"[\\/:*?\"<>|]", "-", artigo.get("titulo", "artigo"))[:70]
    nome_arquivo = f"[JUSBRASIL] {data_str} - {titulo_limpo}.docx"

    pasta_id = (
        os.getenv("JUSBRASIL_DRIVE_FOLDER_ID", "").strip()
        or os.getenv("GOOGLE_PASTA_DESTINO", "").strip()
    )
    if not os.getenv("JUSBRASIL_DRIVE_FOLDER_ID"):
        print("   AVISO: JUSBRASIL_DRIVE_FOLDER_ID nao configurado")
        print("   Usando GOOGLE_PASTA_DESTINO como fallback")

    arquivo_id, link_drive = salvar_no_drive(drive_service, docx_bytes, nome_arquivo, pasta_id)
    print(f"   Arquivo: {nome_arquivo}")
    print(f"   Link: {link_drive}")

    # Historico
    historico.append({
        "data": datetime.now().isoformat(),
        "titulo": artigo.get("titulo", ""),
        "titulo_origem": noticia["titulo"],
        "link_origem": noticia["link"],
        "palavra_chave": artigo.get("palavra_chave_principal", ""),
        "arquivo_drive": arquivo_id,
        "link_drive": link_drive,
    })
    salvar_historico(historico)
    print(f"   Historico: {len(historico)} artigos registrados")

    # WhatsApp
    if not args.sem_notificar:
        print("\nNotificando revisor via WhatsApp...")
        if notificar_whatsapp(artigo, link_drive):
            print("   Notificacao enviada!")
        else:
            print("   Notificacao nao enviada (configure JUSBRASIL_NOTIFICAR_TELEFONE no .env)")

    print("\n" + "=" * 64)
    print("CONCLUIDO! Artigo aguardando revisao no Drive.")
    print("=" * 64)
    return 0


if __name__ == "__main__":
    sys.exit(main())
