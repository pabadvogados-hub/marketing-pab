"""
KIT APRESENTACAO PREMIUM - Upload como Google Slides
Renderiza (se preciso), monta .pptx 16:9 com cada slide em tela cheia e sobe no
Drive CONVERTENDO para Google Slides (apresentacao nativa). Sobe tambem PDF + HTML.

Uso:
    python MARKETING/_KIT_APRESENTACAO/subir_google_slides.py "caminho/deck.html" \
        --nome "Titulo no Drive" --pasta "Marketing"

Regras:
- NAO cria pastas no Drive. Usa a pasta informada se existir; senao, raiz.
"""
import os
import sys
import glob
import argparse
from pathlib import Path

sys.path.insert(0, os.path.abspath(os.path.join(os.path.dirname(__file__), "..", "..")))
sys.path.insert(0, os.path.dirname(__file__))

from pptx import Presentation
from pptx.util import Inches
from INTEGRACOES.google_integration import autenticar_google
from googleapiclient.http import MediaFileUpload
import render as renderer

GSLIDES = "application/vnd.google-apps.presentation"
PPTX_MIME = "application/vnd.openxmlformats-officedocument.presentationml.presentation"


def build_pptx(slides_dir, pptx_path):
    pngs = sorted(glob.glob(os.path.join(slides_dir, "slide_*.png")))
    if not pngs:
        print("ERRO: sem PNGs em", slides_dir); sys.exit(1)
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]
    for png in pngs:
        s = prs.slides.add_slide(blank)
        s.shapes.add_picture(png, 0, 0, width=prs.slide_width, height=prs.slide_height)
    prs.save(pptx_path)
    print(f"PPTX ({len(pngs)} slides): {pptx_path}")
    return pptx_path


def achar_pasta(drive, nome):
    if not nome:
        return None, None
    q = (f"name='{nome}' and mimeType='application/vnd.google-apps.folder' and trashed=false")
    fs = drive.files().list(q=q, fields="files(id,name)").execute().get("files", [])
    return (fs[0]["id"], fs[0]["name"]) if fs else (None, None)


def upload(drive, path, nome, mime, parent, convert_to=None):
    meta = {"name": nome}
    if convert_to:
        meta["mimeType"] = convert_to
    if parent:
        meta["parents"] = [parent]
    media = MediaFileUpload(path, mimetype=mime, resumable=True)
    f = drive.files().create(body=meta, media_body=media, fields="id,name,webViewLink").execute()
    print(f"  OK {f['name']}\n     {f.get('webViewLink')}")
    return f


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("html")
    ap.add_argument("--nome", default=None, help="nome base no Drive")
    ap.add_argument("--pasta", default="Marketing", help="pasta existente no Drive")
    ap.add_argument("--so-slides", action="store_true", help="sobe apenas o Google Slides")
    a = ap.parse_args()

    html = os.path.abspath(a.html)
    stem = Path(html).stem
    base = os.path.dirname(html)
    slides_dir = os.path.join(base, stem + "_slides")
    pdf = os.path.join(base, stem + ".pdf")
    pptx = os.path.join(base, stem + ".pptx")
    nome = a.nome or stem.replace("_", " ")

    if not glob.glob(os.path.join(slides_dir, "slide_*.png")):
        print("PNGs nao encontrados, renderizando...\n")
        renderer.render(html)
    build_pptx(slides_dir, pptx)

    drive, _ = autenticar_google()
    pid, pname = achar_pasta(drive, a.pasta)
    print("Destino:", pname or "Meu Drive (raiz)", "\n")

    upload(drive, pptx, nome, PPTX_MIME, pid, convert_to=GSLIDES)
    if not a.so_slides:
        if os.path.exists(pdf):
            upload(drive, pdf, nome + ".pdf", "application/pdf", pid)
        if os.path.exists(html):
            upload(drive, html, stem + ".html", "text/html", pid)


if __name__ == "__main__":
    main()
